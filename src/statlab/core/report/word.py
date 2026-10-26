"""Informe Word (.docx) editable con python-docx: el mismo contenido que el PDF."""
from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from statlab.core.report.document import Bloque, Figura, Parrafo, SaltoPagina, Tabla, Titulo

MUTED = RGBColor(0x55, 0x54, 0x4F)


def _nota(doc, texto: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(texto)
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = MUTED


def render(bloques: list[Bloque], path: str | Path, titulo: str) -> Path:
    path = Path(path)
    doc = Document()
    doc.core_properties.title = titulo
    doc.core_properties.author = "StatLab"
    section = doc.sections[0]
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(2))
    doc.styles["Normal"].font.size = Pt(10)

    for b in bloques:
        if isinstance(b, Titulo):
            doc.add_heading(b.texto, level=min(b.nivel, 3)).paragraph_format.keep_with_next = True
        elif isinstance(b, Parrafo):
            if b.nota:
                _nota(doc, b.texto)
            else:
                doc.add_paragraph(b.texto)
        elif isinstance(b, Tabla):
            table = doc.add_table(rows=1, cols=len(b.cabeceras))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            for j, h in enumerate(b.cabeceras):
                cell = table.rows[0].cells[j]
                cell.text = ""
                run = cell.paragraphs[0].add_run(str(h))
                run.bold = True
                run.font.size = Pt(8.5)
            for fila in b.filas:
                cells = table.add_row().cells
                for j, v in enumerate(fila):
                    cells[j].text = ""
                    par = cells[j].paragraphs[0]
                    run = par.add_run("" if v is None else str(v))
                    run.font.size = Pt(8.5)
                    if j not in b.izquierda:
                        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if b.pie:
                _nota(doc, b.pie)
        elif isinstance(b, Figura):
            doc.add_picture(io.BytesIO(b.png), width=Cm(min(b.ancho_cm, 17)))
            if b.pie:
                _nota(doc, b.pie)
        elif isinstance(b, SaltoPagina):
            doc.add_page_break()
    doc.save(str(path))
    return path
